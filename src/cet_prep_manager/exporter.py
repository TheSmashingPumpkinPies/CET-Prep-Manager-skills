"""Excel and chart export engine for CET Prep Manager.

Produces a professional multi-sheet Excel workbook with embedded longitudinal
trend charts demonstrating long-term dynamic progress across all CET modules.
"""

from __future__ import annotations

from datetime import datetime, timezone
import io
import json
from pathlib import Path
import sqlite3
from typing import Any

import matplotlib
matplotlib.use("Agg")  # Non-interactive backend for server/CLI environments
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from cet_prep_manager.analytics import (
    build_learner_state_snapshot,
    compute_least_squares_slope,
    compute_moving_average,
    compute_volatility,
)
from cet_prep_manager.db import (
    RecordNotFoundError,
    get_active_learner_profile,
    get_default_db_path,
    get_learner_profile,
    list_error_events,
    list_exam_attempts,
    list_plan_versions,
    list_section_results,
    list_subjective_assessments,
)

# Styling Constants
NAVY_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Segoe UI", size=9)
BOLD_DATA_FONT = Font(name="Segoe UI", size=9, bold=True)
SUBTITLE_FONT = Font(name="Segoe UI", size=12, bold=True, color="1F497D")

THIN_SIDE = Side(border_style="thin", color="D9D9D9")
BORDER_BOX = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _style_header_row(ws: Any, row_idx: int, cols_count: int) -> None:
    """Apply standard navy header styling to a row."""
    for col in range(1, cols_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = NAVY_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_BOX


def _auto_fit_columns(ws: Any, min_col: int = 1, max_col: int = 20) -> None:
    """Adjust column widths based on content length."""
    for col in range(min_col, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                val_str = str(val).split("\n")[0]
                length = len(val_str.encode("utf-8", "ignore"))
                if length > max_len:
                    max_len = length
        calculated = max(max_len + 4, 11)
        ws.column_dimensions[col_letter].width = min(calculated, 42)


# =====================================================================
# Chart Generators (Focus on Long-Term Dynamic Progress)
# =====================================================================


def _render_overall_progress_chart(
    list_pts: list[tuple[str, float]],
    read_pts: list[tuple[str, float]],
    write_pts: list[tuple[str, float]],
    trans_pts: list[tuple[str, float]],
) -> io.BytesIO:
    """Render comprehensive longitudinal multi-panel progress chart."""
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(10, 6.2), dpi=130)
    fig.patch.set_facecolor("#FAFAFA")

    # Panel 1: Objective Sections (Accuracy over time)
    ax1.set_facecolor("#FFFFFF")
    ax1.set_title("Objective Sections: Accuracy & Moving Averages", fontsize=11, fontweight="bold", color="#1F497D", pad=8)
    ax1.axhline(0.60, color="#E67E22", linestyle="--", linewidth=1, alpha=0.7, label="Passing Target (60%)")
    ax1.axhline(0.80, color="#27AE60", linestyle="--", linewidth=1, alpha=0.7, label="Excellence Target (80%)")

    if list_pts:
        x_l = list(range(1, len(list_pts) + 1))
        y_l = [p[1] for p in list_pts]
        ax1.plot(x_l, y_l, marker="o", markersize=5, color="#2980B9", linewidth=1.5, label="Listening Observed")
        if len(y_l) >= 3:
            ma3_l = [sum(y_l[max(0, i - 2) : i + 1]) / len(y_l[max(0, i - 2) : i + 1]) for i in range(len(y_l))]
            ax1.plot(x_l, ma3_l, linestyle="--", color="#1A5276", linewidth=2, label="Listening MA3")

    if read_pts:
        x_r = list(range(1, len(read_pts) + 1))
        y_r = [p[1] for p in read_pts]
        ax1.plot(x_r, y_r, marker="s", markersize=5, color="#8E44AD", linewidth=1.5, label="Reading Observed")
        if len(y_r) >= 3:
            ma3_r = [sum(y_r[max(0, i - 2) : i + 1]) / len(y_r[max(0, i - 2) : i + 1]) for i in range(len(y_r))]
            ax1.plot(x_r, ma3_r, linestyle="--", color="#5B2C6F", linewidth=2, label="Reading MA3")

    ax1.set_ylabel("Accuracy", fontsize=9, fontweight="bold")
    ax1.set_ylim(0.2, 1.05)
    ax1.grid(True, linestyle=":", alpha=0.6)
    ax1.legend(loc="lower right", fontsize=8, framealpha=0.9)

    # Panel 2: Subjective Sections (Writing & Translation estimated scores)
    ax2.set_facecolor("#FFFFFF")
    ax2.set_title("Subjective Sections: Official-Aligned Estimates (1–15 Scale)", fontsize=11, fontweight="bold", color="#1F497D", pad=8)
    ax2.axhline(11, color="#27AE60", linestyle="--", linewidth=1, alpha=0.6, label="Band 11 Target (10-12)")
    ax2.axhline(8, color="#E67E22", linestyle="--", linewidth=1, alpha=0.6, label="Band 8 Base (7-9)")

    if write_pts:
        x_w = list(range(1, len(write_pts) + 1))
        y_w = [p[1] for p in write_pts]
        ax2.plot(x_w, y_w, marker="^", markersize=6, color="#D35400", linewidth=1.8, label="Writing Est.")
        if len(y_w) >= 3:
            ma3_w = [sum(y_w[max(0, i - 2) : i + 1]) / len(y_w[max(0, i - 2) : i + 1]) for i in range(len(y_w))]
            ax2.plot(x_w, ma3_w, linestyle=":", color="#BA4A00", linewidth=2, label="Writing MA3")

    if trans_pts:
        x_t = list(range(1, len(trans_pts) + 1))
        y_t = [p[1] for p in trans_pts]
        ax2.plot(x_t, y_t, marker="d", markersize=6, color="#16A085", linewidth=1.8, label="Translation Est.")
        if len(y_t) >= 3:
            ma3_t = [sum(y_t[max(0, i - 2) : i + 1]) / len(y_t[max(0, i - 2) : i + 1]) for i in range(len(y_t))]
            ax2.plot(x_t, ma3_t, linestyle=":", color="#117864", linewidth=2, label="Translation MA3")

    ax2.set_xlabel("Attempt Sequence Index", fontsize=9, fontweight="bold")
    ax2.set_ylabel("Score (1–15)", fontsize=9, fontweight="bold")
    ax2.set_ylim(2, 15.5)
    ax2.grid(True, linestyle=":", alpha=0.6)
    ax2.legend(loc="lower right", fontsize=8, framealpha=0.9)

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _render_objective_trend_chart(title: str, points: list[tuple[str, float]], color_theme: str) -> io.BytesIO:
    """Render single objective section accuracy trend with MA3, MA5 and linear trend."""
    fig, ax = plt.subplots(figsize=(8.5, 4.6), dpi=130)
    fig.patch.set_facecolor("#FAFAFA")
    ax.set_facecolor("#FFFFFF")
    ax.set_title(f"{title}: Longitudinal Accuracy & Trend Dynamics", fontsize=11, fontweight="bold", color="#1F497D", pad=10)

    ax.axhline(0.60, color="#E67E22", linestyle="--", linewidth=1, alpha=0.7, label="Passing Target (60%)")
    ax.axhline(0.80, color="#27AE60", linestyle="--", linewidth=1, alpha=0.7, label="Target (80%)")

    if points:
        x = list(range(1, len(points) + 1))
        y = [p[1] for p in points]
        ax.plot(x, y, marker="o", markersize=6, color=color_theme, linewidth=1.6, label="Observed Accuracy")

        for xi, yi in zip(x, y):
            ax.annotate(f"{yi*100:.0f}%", (xi, yi), textcoords="offset points", xytext=(0, 6), ha="center", fontsize=7.5)

        if len(y) >= 3:
            ma3 = [sum(y[max(0, i - 2) : i + 1]) / len(y[max(0, i - 2) : i + 1]) for i in range(len(y))]
            ax.plot(x, ma3, linestyle="--", color="#E67E22", linewidth=2.2, label="Short MA (MA3)")

        if len(y) >= 5:
            ma5 = [sum(y[max(0, i - 4) : i + 1]) / len(y[max(0, i - 4) : i + 1]) for i in range(len(y))]
            ax.plot(x, ma5, linestyle="-.", color="#27AE60", linewidth=2.2, label="Long MA (MA5)")

        if len(y) >= 3:
            slope = compute_least_squares_slope(y, max_window=len(y))
            if slope is not None:
                x_mean = sum(x) / len(x)
                y_mean = sum(y) / len(y)
                trend_y = [y_mean + slope * (xi - x_mean) for xi in x]
                direction = "Progressing" if slope > 0 else ("Declining" if slope < 0 else "Stable")
                ax.plot(x, trend_y, linestyle=":", color="#7F8C8D", linewidth=1.8, label=f"Trend Slope: {slope:+.3f} ({direction})")

    ax.set_xlabel("Attempt Sequence Index", fontsize=9, fontweight="bold")
    ax.set_ylabel("Accuracy", fontsize=9, fontweight="bold")
    ax.set_ylim(0.25, 1.05)
    ax.grid(True, linestyle=":", alpha=0.6)
    ax.legend(loc="lower right", fontsize=8.5, framealpha=0.9)

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _render_subjective_band_chart(title: str, assessments: list[dict[str, Any]]) -> io.BytesIO:
    """Render subjective assessment longitudinal progress with uncertainty range and anchor bands."""
    fig, ax = plt.subplots(figsize=(8.8, 4.8), dpi=130)
    fig.patch.set_facecolor("#FAFAFA")
    ax.set_facecolor("#FFFFFF")
    ax.set_title(f"{title}: Longitudinal Score Bands & Uncertainty Range", fontsize=11, fontweight="bold", color="#1F497D", pad=10)

    # Shaded anchor band tiers (official CET bands)
    bands = [
        (13, 15, "#D5F5E3", "Band 14 (13–15 Excellent)"),
        (10, 12, "#E8F8F5", "Band 11 (10–12 Good)"),
        (7, 9, "#FEF9E7", "Band 8 (7–9 Pass)"),
        (4, 6, "#FDEDEC", "Band 5 (4–6 Weak)"),
        (1, 3, "#FADBD8", "Band 2 (1–3 Poor)"),
    ]
    for low, high, bg_color, _ in bands:
        ax.axhspan(low - 0.45, high + 0.45, color=bg_color, alpha=0.4)

    if assessments:
        x = list(range(1, len(assessments) + 1))
        scores = [float(a["estimated_score"]) for a in assessments]
        lows = [float(a.get("score_low", a["estimated_score"])) for a in assessments]
        highs = [float(a.get("score_high", a["estimated_score"])) for a in assessments]

        # Uncertainty envelope fill
        ax.fill_between(x, lows, highs, color="#3498DB", alpha=0.25, label="Likely Range [Score Low - High]")
        ax.plot(x, scores, marker="o", markersize=6.5, color="#2980B9", linewidth=2.0, label="Estimated Point Score")

        for xi, sc in zip(x, scores):
            ax.annotate(f"{sc:.0f}/15", (xi, sc), textcoords="offset points", xytext=(0, 6), ha="center", fontsize=8, fontweight="bold")

        if len(scores) >= 3:
            ma3 = [sum(scores[max(0, i - 2) : i + 1]) / len(scores[max(0, i - 2) : i + 1]) for i in range(len(scores))]
            ax.plot(x, ma3, linestyle="--", color="#E67E22", linewidth=2.2, label="Moving Average (MA3)")

        ax.legend(loc="lower right", fontsize=8.5, framealpha=0.9)
    else:
        ax.text(0.5, 0.5, f"No {title} Assessments Recorded Yet", ha="center", va="center", fontsize=11, color="#7F8C8D")

    ax.set_xlabel("Attempt Sequence Index", fontsize=9, fontweight="bold")
    ax.set_ylabel("Official-Aligned Estimate (1–15)", fontsize=9, fontweight="bold")
    ax.set_ylim(1, 15.8)
    ax.grid(True, linestyle=":", alpha=0.6)

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _render_error_distribution_chart(errors: list[dict[str, Any]]) -> io.BytesIO:
    """Render horizontal stacked bar chart showing error frequency by category and resolution state."""
    fig, ax = plt.subplots(figsize=(8.5, 4.6), dpi=130)
    fig.patch.set_facecolor("#FAFAFA")
    ax.set_facecolor("#FFFFFF")
    ax.set_title("Top Recurring Error Codes & Resolution Progression", fontsize=11, fontweight="bold", color="#1F497D", pad=10)

    # Tally counts by taxonomy_code and state
    stats: dict[str, dict[str, int]] = {}
    for err in errors:
        code = err.get("taxonomy_code", "UNKNOWN")
        st = err.get("resolved_state", "new")
        if code not in stats:
            stats[code] = {"new": 0, "recurrent": 0, "improving": 0, "resolved": 0}
        stats[code][st] = stats[code].get(st, 0) + 1

    if not stats:
        ax.text(0.5, 0.5, "No Error Events Recorded Yet", ha="center", va="center", fontsize=11, color="#7F8C8D")
        ax.set_axis_off()
    else:
        # Sort by total occurrences desc, take top 8
        sorted_items = sorted(stats.items(), key=lambda item: sum(item[1].values()), reverse=True)[:8]
        sorted_items.reverse()  # For horizontal bar display from top to bottom
        categories = [item[0] for item in sorted_items]

        new_counts = [item[1]["new"] for item in sorted_items]
        recurrent_counts = [item[1]["recurrent"] for item in sorted_items]
        improving_counts = [item[1]["improving"] for item in sorted_items]
        resolved_counts = [item[1]["resolved"] for item in sorted_items]

        ax.barh(categories, new_counts, label="New", color="#E74C3C", alpha=0.85)
        ax.barh(categories, recurrent_counts, left=new_counts, label="Recurrent", color="#E67E22", alpha=0.85)

        left_imp = [n + r for n, r in zip(new_counts, recurrent_counts)]
        ax.barh(categories, improving_counts, left=left_imp, label="Improving", color="#F1C40F", alpha=0.85)

        left_res = [l + i for l, i in zip(left_imp, improving_counts)]
        ax.barh(categories, resolved_counts, left=left_res, label="Resolved", color="#2ECC71", alpha=0.85)

        ax.set_xlabel("Occurrences Count", fontsize=9, fontweight="bold")
        ax.grid(True, axis="x", linestyle=":", alpha=0.6)
        ax.legend(loc="lower right", fontsize=8.5, framealpha=0.9)

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


# =====================================================================
# Main Workbook Exporter
# =====================================================================


def generate_excel_report(
    conn: sqlite3.Connection,
    learner_id: str,
    output_path: str | Path,
) -> Path:
    """Generate professional 9-sheet Excel workbook with embedded longitudinal trend charts."""
    profile = get_learner_profile(conn, learner_id)
    if not profile:
        raise RecordNotFoundError(f"Learner profile '{learner_id}' not found")

    target_path = Path(output_path).expanduser().resolve()
    target_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    snapshot = build_learner_state_snapshot(conn, learner_id)
    attempts = list_exam_attempts(conn, learner_id=learner_id, limit=500)
    attempts.sort(key=lambda a: a.get("attempted_at", ""))  # Chronological

    all_sections = list_section_results(conn, learner_id=learner_id)
    writing_assessments = list_subjective_assessments(conn, learner_id=learner_id, section="writing", limit=500)
    translation_assessments = list_subjective_assessments(conn, learner_id=learner_id, section="translation", limit=500)
    error_events = list_error_events(conn, learner_id=learner_id)
    plan_versions = list_plan_versions(conn, learner_id=learner_id)

    # 1. Sheet: Overview
    ws_ov = wb.create_sheet("Overview")
    ws_ov.views.sheetView[0].showGridLines = True
    ws_ov.freeze_panes = "A2"

    ws_ov["A1"] = "CET PREPARATION DASHBOARD & LONGITUDINAL SUMMARY"
    ws_ov["A1"].font = SUBTITLE_FONT

    ws_ov["A3"] = "Learner ID:"
    ws_ov["B3"] = profile["learner_id"]
    ws_ov["A4"] = "Target Exam:"
    ws_ov["B4"] = profile["target_exam"]
    ws_ov["A5"] = "Target Date:"
    ws_ov["B5"] = profile.get("target_exam_date") or "Not set"
    ws_ov["A6"] = "Target Score:"
    ws_ov["B6"] = profile.get("target_reported_score") or "Not set"
    ws_ov["A7"] = "Generated At:"
    ws_ov["B7"] = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    for r in range(3, 8):
        ws_ov.cell(row=r, column=1).font = BOLD_DATA_FONT
        ws_ov.cell(row=r, column=2).font = DATA_FONT

    # Module summary table
    ov_headers = ["Module", "Samples (N)", "Latest Metric", "MA3", "MA5", "Trend Slope", "Volatility", "Sufficiency"]
    for c_idx, h in enumerate(ov_headers, start=1):
        ws_ov.cell(row=9, column=c_idx, value=h)
    _style_header_row(ws_ov, 9, len(ov_headers))

    row_pos = 10
    for mod_name in ["listening", "reading", "writing", "translation"]:
        st = snapshot.modules.get(mod_name)
        if not st:
            continue
        is_subj = mod_name in ("writing", "translation")
        ws_ov.cell(row=row_pos, column=1, value=mod_name.capitalize()).font = BOLD_DATA_FONT
        ws_ov.cell(row=row_pos, column=2, value=st.n).alignment = ALIGN_CENTER

        c_latest = ws_ov.cell(row=row_pos, column=3)
        if st.latest is not None:
            c_latest.value = st.latest
            c_latest.number_format = "0.0" if is_subj else "0.0%"
        else:
            c_latest.value = "N/A"
        c_latest.alignment = ALIGN_RIGHT

        c_ma3 = ws_ov.cell(row=row_pos, column=4)
        if st.ma3 is not None:
            c_ma3.value = st.ma3
            c_ma3.number_format = "0.0" if is_subj else "0.0%"
        else:
            c_ma3.value = "N/A"
        c_ma3.alignment = ALIGN_RIGHT

        c_ma5 = ws_ov.cell(row=row_pos, column=5)
        if st.ma5 is not None:
            c_ma5.value = st.ma5
            c_ma5.number_format = "0.0" if is_subj else "0.0%"
        else:
            c_ma5.value = "N/A"
        c_ma5.alignment = ALIGN_RIGHT

        c_slope = ws_ov.cell(row=row_pos, column=6)
        if st.trend_slope is not None:
            c_slope.value = st.trend_slope
            c_slope.number_format = "+0.0000;-0.0000;0.0000"
        else:
            c_slope.value = "N/A"
        c_slope.alignment = ALIGN_RIGHT

        c_vol = ws_ov.cell(row=row_pos, column=7)
        if st.volatility is not None:
            c_vol.value = st.volatility
            c_vol.number_format = "0.0000"
        else:
            c_vol.value = "N/A"
        c_vol.alignment = ALIGN_RIGHT

        c_suff = ws_ov.cell(row=row_pos, column=8, value=st.sufficiency.value)
        c_suff.alignment = ALIGN_CENTER
        c_suff.font = DATA_FONT

        for c in range(1, 9):
            ws_ov.cell(row=row_pos, column=c).border = BORDER_BOX
        row_pos += 1

    # Prepare data points for charts
    list_series = [(r.get("attempted_at", ""), float(r["accuracy"])) for r in all_sections if r.get("section") == "listening" and r.get("accuracy") is not None]
    read_series = [(r.get("attempted_at", ""), float(r["accuracy"])) for r in all_sections if r.get("section") == "reading" and r.get("accuracy") is not None]
    write_series = [(a.get("created_at", ""), float(a["estimated_score"])) for a in writing_assessments if a.get("estimated_score") is not None]
    trans_series = [(a.get("created_at", ""), float(a["estimated_score"])) for a in translation_assessments if a.get("estimated_score") is not None]

    # Embed overall progress chart into Overview
    chart_buf_ov = _render_overall_progress_chart(list_series, read_series, write_series, trans_series)
    img_ov = OpenpyxlImage(chart_buf_ov)
    img_ov.width = 660
    img_ov.height = 410
    ws_ov.add_image(img_ov, "J3")

    _auto_fit_columns(ws_ov, 1, 8)

    # 2. Sheet: Mock History
    ws_mh = wb.create_sheet("Mock History")
    ws_mh.freeze_panes = "A2"
    mh_headers = [
        "Attempt ID",
        "Date",
        "Level",
        "Type",
        "Source Key",
        "Source Kind",
        "Official Reported Score",
        "Duration (s)",
        "Notes",
    ]
    for c_idx, h in enumerate(mh_headers, start=1):
        ws_mh.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_mh, 1, len(mh_headers))

    for r_idx, att in enumerate(attempts, start=2):
        ws_mh.cell(row=r_idx, column=1, value=att["attempt_id"]).alignment = ALIGN_CENTER
        ws_mh.cell(row=r_idx, column=2, value=att.get("attempted_at", "")[:10]).alignment = ALIGN_CENTER
        ws_mh.cell(row=r_idx, column=3, value=att["exam_level"]).alignment = ALIGN_CENTER
        ws_mh.cell(row=r_idx, column=4, value=att["attempt_type"]).alignment = ALIGN_CENTER
        ws_mh.cell(row=r_idx, column=5, value=att.get("source_key") or "N/A")
        ws_mh.cell(row=r_idx, column=6, value=att.get("source_kind") or "N/A").alignment = ALIGN_CENTER
        ws_mh.cell(
            row=r_idx,
            column=7,
            value=att.get("official_reported_score"),
        ).alignment = ALIGN_RIGHT
        ws_mh.cell(
            row=r_idx,
            column=8,
            value=att.get("duration_seconds"),
        ).alignment = ALIGN_RIGHT
        ws_mh.cell(row=r_idx, column=9, value=att.get("notes") or "")
        for c in range(1, len(mh_headers) + 1):
            ws_mh.cell(row=r_idx, column=c).font = DATA_FONT
            ws_mh.cell(row=r_idx, column=c).border = BORDER_BOX
    _auto_fit_columns(ws_mh, 1, len(mh_headers))

    # 3. Sheet: Listening
    ws_li = wb.create_sheet("Listening")
    ws_li.freeze_panes = "A2"
    li_headers = ["Date", "Subtype", "Correct", "Total", "Accuracy", "Duration (s)", "Practice Index", "Provenance"]
    for c_idx, h in enumerate(li_headers, start=1):
        ws_li.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_li, 1, len(li_headers))

    li_rows = [r for r in all_sections if r.get("section") == "listening"]
    for r_idx, sec in enumerate(li_rows, start=2):
        ws_li.cell(row=r_idx, column=1, value=sec.get("attempted_at", "")[:10]).alignment = ALIGN_CENTER
        ws_li.cell(row=r_idx, column=2, value=sec.get("subtype") or "all").alignment = ALIGN_CENTER
        ws_li.cell(row=r_idx, column=3, value=sec.get("correct_count")).alignment = ALIGN_RIGHT
        ws_li.cell(row=r_idx, column=4, value=sec.get("total_count")).alignment = ALIGN_RIGHT
        c_acc = ws_li.cell(row=r_idx, column=5, value=sec.get("accuracy"))
        if sec.get("accuracy") is not None:
            c_acc.number_format = "0.0%"
        c_acc.alignment = ALIGN_RIGHT
        ws_li.cell(row=r_idx, column=6, value=sec.get("duration_seconds")).alignment = ALIGN_RIGHT
        ws_li.cell(row=r_idx, column=7, value=sec.get("practice_index")).alignment = ALIGN_RIGHT
        ws_li.cell(row=r_idx, column=8, value=sec.get("provenance") or "recorded").alignment = ALIGN_CENTER
        for c in range(1, len(li_headers) + 1):
            ws_li.cell(row=r_idx, column=c).font = DATA_FONT
            ws_li.cell(row=r_idx, column=c).border = BORDER_BOX

    # Embed Listening Chart
    chart_buf_li = _render_objective_trend_chart("Listening", list_series, "#2980B9")
    img_li = OpenpyxlImage(chart_buf_li)
    img_li.width = 560
    img_li.height = 300
    ws_li.add_image(img_li, "J2")
    _auto_fit_columns(ws_li, 1, len(li_headers))

    # 4. Sheet: Reading
    ws_re = wb.create_sheet("Reading")
    ws_re.freeze_panes = "A2"
    re_headers = ["Date", "Subtype", "Correct", "Total", "Accuracy", "Duration (s)", "Practice Index", "Provenance"]
    for c_idx, h in enumerate(re_headers, start=1):
        ws_re.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_re, 1, len(re_headers))

    re_rows = [r for r in all_sections if r.get("section") == "reading"]
    for r_idx, sec in enumerate(re_rows, start=2):
        ws_re.cell(row=r_idx, column=1, value=sec.get("attempted_at", "")[:10]).alignment = ALIGN_CENTER
        ws_re.cell(row=r_idx, column=2, value=sec.get("subtype") or "all").alignment = ALIGN_CENTER
        ws_re.cell(row=r_idx, column=3, value=sec.get("correct_count")).alignment = ALIGN_RIGHT
        ws_re.cell(row=r_idx, column=4, value=sec.get("total_count")).alignment = ALIGN_RIGHT
        c_acc = ws_re.cell(row=r_idx, column=5, value=sec.get("accuracy"))
        if sec.get("accuracy") is not None:
            c_acc.number_format = "0.0%"
        c_acc.alignment = ALIGN_RIGHT
        ws_re.cell(row=r_idx, column=6, value=sec.get("duration_seconds")).alignment = ALIGN_RIGHT
        ws_re.cell(row=r_idx, column=7, value=sec.get("practice_index")).alignment = ALIGN_RIGHT
        ws_re.cell(row=r_idx, column=8, value=sec.get("provenance") or "recorded").alignment = ALIGN_CENTER
        for c in range(1, len(re_headers) + 1):
            ws_re.cell(row=r_idx, column=c).font = DATA_FONT
            ws_re.cell(row=r_idx, column=c).border = BORDER_BOX

    # Embed Reading Chart
    chart_buf_re = _render_objective_trend_chart("Reading", read_series, "#8E44AD")
    img_re = OpenpyxlImage(chart_buf_re)
    img_re.width = 560
    img_re.height = 300
    ws_re.add_image(img_re, "J2")
    _auto_fit_columns(ws_re, 1, len(re_headers))

    # 5. Sheet: Writing
    ws_wr = wb.create_sheet("Writing")
    ws_wr.freeze_panes = "A2"
    wr_headers = ["Date", "Score (1-15)", "Low", "High", "Anchor", "Band", "Confidence", "Assessor", "Rubric", "Rationale"]
    for c_idx, h in enumerate(wr_headers, start=1):
        ws_wr.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_wr, 1, len(wr_headers))

    for r_idx, w in enumerate(writing_assessments, start=2):
        ws_wr.cell(row=r_idx, column=1, value=w.get("created_at", "")[:10]).alignment = ALIGN_CENTER
        ws_wr.cell(row=r_idx, column=2, value=w.get("estimated_score")).alignment = ALIGN_RIGHT
        ws_wr.cell(row=r_idx, column=3, value=w.get("score_low")).alignment = ALIGN_RIGHT
        ws_wr.cell(row=r_idx, column=4, value=w.get("score_high")).alignment = ALIGN_RIGHT
        ws_wr.cell(row=r_idx, column=5, value=w.get("anchor_band")).alignment = ALIGN_CENTER
        ws_wr.cell(row=r_idx, column=6, value=f"{w.get('band_low')}–{w.get('band_high')}").alignment = ALIGN_CENTER
        ws_wr.cell(row=r_idx, column=7, value=w.get("confidence")).alignment = ALIGN_CENTER
        ws_wr.cell(row=r_idx, column=8, value=w.get("assessor_model")).alignment = ALIGN_LEFT
        ws_wr.cell(row=r_idx, column=9, value=w.get("rubric_version")).alignment = ALIGN_CENTER
        ws_wr.cell(row=r_idx, column=10, value=(w.get("rationale_md") or "")[:80])
        for c in range(1, len(wr_headers) + 1):
            ws_wr.cell(row=r_idx, column=c).font = DATA_FONT
            ws_wr.cell(row=r_idx, column=c).border = BORDER_BOX

    # Embed Writing Chart
    chart_buf_wr = _render_subjective_band_chart("Writing", writing_assessments)
    img_wr = OpenpyxlImage(chart_buf_wr)
    img_wr.width = 580
    img_wr.height = 310
    ws_wr.add_image(img_wr, "L2")
    _auto_fit_columns(ws_wr, 1, len(wr_headers))

    # 6. Sheet: Translation
    ws_tr = wb.create_sheet("Translation")
    ws_tr.freeze_panes = "A2"
    tr_headers = ["Date", "Score (1-15)", "Low", "High", "Anchor", "Band", "Confidence", "Assessor", "Rubric", "Rationale"]
    for c_idx, h in enumerate(tr_headers, start=1):
        ws_tr.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_tr, 1, len(tr_headers))

    for r_idx, t in enumerate(translation_assessments, start=2):
        ws_tr.cell(row=r_idx, column=1, value=t.get("created_at", "")[:10]).alignment = ALIGN_CENTER
        ws_tr.cell(row=r_idx, column=2, value=t.get("estimated_score")).alignment = ALIGN_RIGHT
        ws_tr.cell(row=r_idx, column=3, value=t.get("score_low")).alignment = ALIGN_RIGHT
        ws_tr.cell(row=r_idx, column=4, value=t.get("score_high")).alignment = ALIGN_RIGHT
        ws_tr.cell(row=r_idx, column=5, value=t.get("anchor_band")).alignment = ALIGN_CENTER
        ws_tr.cell(row=r_idx, column=6, value=f"{t.get('band_low')}–{t.get('band_high')}").alignment = ALIGN_CENTER
        ws_tr.cell(row=r_idx, column=7, value=t.get("confidence")).alignment = ALIGN_CENTER
        ws_tr.cell(row=r_idx, column=8, value=t.get("assessor_model")).alignment = ALIGN_LEFT
        ws_tr.cell(row=r_idx, column=9, value=t.get("rubric_version")).alignment = ALIGN_CENTER
        ws_tr.cell(row=r_idx, column=10, value=(t.get("rationale_md") or "")[:80])
        for c in range(1, len(tr_headers) + 1):
            ws_tr.cell(row=r_idx, column=c).font = DATA_FONT
            ws_tr.cell(row=r_idx, column=c).border = BORDER_BOX

    # Embed Translation Chart
    chart_buf_tr = _render_subjective_band_chart("Translation", translation_assessments)
    img_tr = OpenpyxlImage(chart_buf_tr)
    img_tr.width = 580
    img_tr.height = 310
    ws_tr.add_image(img_tr, "L2")
    _auto_fit_columns(ws_tr, 1, len(tr_headers))

    # 7. Sheet: Errors
    ws_er = wb.create_sheet("Errors")
    ws_er.freeze_panes = "A2"
    er_headers = ["Error ID", "Date", "Section", "Subtype", "Taxonomy Code", "Severity", "State", "Evidence Note"]
    for c_idx, h in enumerate(er_headers, start=1):
        ws_er.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_er, 1, len(er_headers))

    for r_idx, err in enumerate(error_events, start=2):
        ws_er.cell(row=r_idx, column=1, value=err["error_id"][:8]).alignment = ALIGN_CENTER
        ws_er.cell(row=r_idx, column=2, value=err.get("created_at", "")[:10]).alignment = ALIGN_CENTER
        ws_er.cell(row=r_idx, column=3, value=err["section"]).alignment = ALIGN_CENTER
        ws_er.cell(row=r_idx, column=4, value=err.get("subtype") or "-").alignment = ALIGN_CENTER
        ws_er.cell(row=r_idx, column=5, value=err["taxonomy_code"]).font = BOLD_DATA_FONT
        ws_er.cell(row=r_idx, column=6, value=err["severity"]).alignment = ALIGN_CENTER
        ws_er.cell(row=r_idx, column=7, value=err["resolved_state"]).alignment = ALIGN_CENTER
        ws_er.cell(row=r_idx, column=8, value=err.get("evidence_note") or "")
        for c in range(1, len(er_headers) + 1):
            if c != 5:
                ws_er.cell(row=r_idx, column=c).font = DATA_FONT
            ws_er.cell(row=r_idx, column=c).border = BORDER_BOX

    # Embed Error Stacked Distribution Chart
    chart_buf_er = _render_error_distribution_chart(error_events)
    img_er = OpenpyxlImage(chart_buf_er)
    img_er.width = 560
    img_er.height = 300
    ws_er.add_image(img_er, "J2")
    _auto_fit_columns(ws_er, 1, len(er_headers))

    # 8. Sheet: Plan History
    ws_pl = wb.create_sheet("Plan History")
    ws_pl.freeze_panes = "A2"
    pl_headers = ["Plan ID", "Version", "Valid From", "Valid To", "Primary Priorities", "Rationale", "Module", "Activity", "Target", "Status"]
    for c_idx, h in enumerate(pl_headers, start=1):
        ws_pl.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_pl, 1, len(pl_headers))

    row_pos_pl = 2
    for plan in plan_versions:
        items = plan.get("items", [])
        p_str = json.dumps(plan.get("primary_priorities", []), ensure_ascii=False) if isinstance(plan.get("primary_priorities"), (list, dict)) else str(plan.get("primary_priorities", ""))
        if not items:
            ws_pl.cell(row=row_pos_pl, column=1, value=plan["plan_id"][:8]).alignment = ALIGN_CENTER
            ws_pl.cell(row=row_pos_pl, column=2, value=f"v{plan['version']}").alignment = ALIGN_CENTER
            ws_pl.cell(row=row_pos_pl, column=3, value=plan.get("valid_from", "")[:10]).alignment = ALIGN_CENTER
            ws_pl.cell(row=row_pos_pl, column=4, value=plan.get("valid_to", "")[:10] if plan.get("valid_to") else "-").alignment = ALIGN_CENTER
            ws_pl.cell(row=row_pos_pl, column=5, value=p_str)
            ws_pl.cell(row=row_pos_pl, column=6, value=plan.get("rationale_md") or "-")
            ws_pl.cell(row=row_pos_pl, column=7, value="-")
            ws_pl.cell(row=row_pos_pl, column=8, value="-")
            ws_pl.cell(row=row_pos_pl, column=9, value="-")
            ws_pl.cell(row=row_pos_pl, column=10, value="-")
            for c in range(1, 11):
                ws_pl.cell(row=row_pos_pl, column=c).font = DATA_FONT
                ws_pl.cell(row=row_pos_pl, column=c).border = BORDER_BOX
            row_pos_pl += 1
        else:
            for it in items:
                ws_pl.cell(row=row_pos_pl, column=1, value=plan["plan_id"][:8]).alignment = ALIGN_CENTER
                ws_pl.cell(row=row_pos_pl, column=2, value=f"v{plan['version']}").alignment = ALIGN_CENTER
                ws_pl.cell(row=row_pos_pl, column=3, value=plan.get("valid_from", "")[:10]).alignment = ALIGN_CENTER
                ws_pl.cell(row=row_pos_pl, column=4, value=plan.get("valid_to", "")[:10] if plan.get("valid_to") else "-").alignment = ALIGN_CENTER
                ws_pl.cell(row=row_pos_pl, column=5, value=p_str)
                ws_pl.cell(row=row_pos_pl, column=6, value=plan.get("rationale_md") or "-")
                ws_pl.cell(row=row_pos_pl, column=7, value=it.get("module")).alignment = ALIGN_CENTER
                ws_pl.cell(row=row_pos_pl, column=8, value=it.get("activity_type")).alignment = ALIGN_LEFT
                tgt = f"{it.get('target_minutes')}m" if it.get("target_minutes") else (f"{it.get('target_count')} items" if it.get("target_count") else "-")
                ws_pl.cell(row=row_pos_pl, column=9, value=tgt).alignment = ALIGN_RIGHT
                ws_pl.cell(row=row_pos_pl, column=10, value=it.get("status")).alignment = ALIGN_CENTER
                for c in range(1, 11):
                    ws_pl.cell(row=row_pos_pl, column=c).font = DATA_FONT
                    ws_pl.cell(row=row_pos_pl, column=c).border = BORDER_BOX
                row_pos_pl += 1
    _auto_fit_columns(ws_pl, 1, len(pl_headers))

    # 9. Sheet: Weekly Summary
    ws_ws = wb.create_sheet("Weekly Summary")
    ws_ws.freeze_panes = "A2"
    ws_headers = ["Reporting Date", "Total Attempts", "Listening Avg", "Reading Avg", "Writing Avg (1-15)", "Translation Avg (1-15)", "Active Errors"]
    for c_idx, h in enumerate(ws_headers, start=1):
        ws_ws.cell(row=1, column=c_idx, value=h)
    _style_header_row(ws_ws, 1, len(ws_headers))

    # Provide summary line
    ws_ws.cell(row=2, column=1, value=datetime.now(timezone.utc).strftime("%Y-%m-%d")).alignment = ALIGN_CENTER
    ws_ws.cell(row=2, column=2, value=len(attempts)).alignment = ALIGN_RIGHT

    l_accs = [p[1] for p in list_series]
    c_la = ws_ws.cell(row=2, column=3)
    if l_accs:
        c_la.value = round(sum(l_accs) / len(l_accs), 4)
        c_la.number_format = "0.0%"
    else:
        c_la.value = "N/A"
    c_la.alignment = ALIGN_RIGHT

    r_accs = [p[1] for p in read_series]
    c_ra = ws_ws.cell(row=2, column=4)
    if r_accs:
        c_ra.value = round(sum(r_accs) / len(r_accs), 4)
        c_ra.number_format = "0.0%"
    else:
        c_ra.value = "N/A"
    c_ra.alignment = ALIGN_RIGHT

    w_scs = [p[1] for p in write_series]
    c_wa = ws_ws.cell(row=2, column=5)
    c_wa.value = round(sum(w_scs) / len(w_scs), 1) if w_scs else "N/A"
    c_wa.alignment = ALIGN_RIGHT

    t_scs = [p[1] for p in trans_series]
    c_ta = ws_ws.cell(row=2, column=6)
    c_ta.value = round(sum(t_scs) / len(t_scs), 1) if t_scs else "N/A"
    c_ta.alignment = ALIGN_RIGHT

    active_errs_count = sum(1 for e in error_events if e.get("resolved_state") in ("new", "recurrent", "improving"))
    ws_ws.cell(row=2, column=7, value=active_errs_count).alignment = ALIGN_RIGHT

    for c in range(1, len(ws_headers) + 1):
        ws_ws.cell(row=2, column=c).font = DATA_FONT
        ws_ws.cell(row=2, column=c).border = BORDER_BOX
    _auto_fit_columns(ws_ws, 1, len(ws_headers))

    # Save workbook
    wb.save(str(target_path))
    return target_path

